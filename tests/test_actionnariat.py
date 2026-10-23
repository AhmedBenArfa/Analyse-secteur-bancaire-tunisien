# Auteur : Ahmed Ben Arfa
"""Tests du fichier d'actionnariat donnees_reference/banques.xlsx (lu par Power BI, requête
Banques_Actionnariat) : périodes continues pour chaque banque de 2019 à aujourd'hui, catégories de la
règle (spécification 5.3), source et page sur chaque ligne."""
import datetime as dt
from pathlib import Path
from openpyxl import load_workbook

FICHIER = Path(__file__).resolve().parent.parent / "donnees_reference" / "banques.xlsx"
BANQUES = {"AB", "ATB", "TJARI", "BH", "BIAT", "BNA", "BT", "BTE", "STB", "UBCI", "UIB", "WIB"}
CATEGORIES = {"public", "groupe étranger", "mixte", "privé sans contrôle étranger"}
DEBUT = dt.date(2019, 1, 1)


def lignes():
    feuille = load_workbook(FICHIER, data_only=True)["Banques"]
    rangees = list(feuille.iter_rows(values_only=True))
    entetes = rangees[0]
    jour = lambda v: v.date() if isinstance(v, dt.datetime) else v
    return [{k: jour(v) for k, v in zip(entetes, r)} for r in rangees[1:] if any(r)]


def test_periodes_continues_2019_a_aujourdhui():
    """Pour chaque banque, les périodes couvrent sans trou ni chevauchement du 01/01/2019 à aujourd'hui :
    la première commence avant 2019 (date_debut vide ou au plus tard le 01/01/2019), la dernière n'a pas de
    fin, et chaque période commence le lendemain de la fin de la précédente (BTE : 30/06/2024 puis
    01/07/2024 ; UBCI : 15/03/2021 puis 16/03/2021 ; Wifak : 31/12/2021 puis 01/01/2022)."""
    par_banque = {}
    for l in lignes():
        par_banque.setdefault(l["code_banque"], []).append(l)
    assert set(par_banque) == BANQUES
    for banque, periodes in par_banque.items():
        periodes.sort(key=lambda l: l["date_debut"] or dt.date.min)
        assert periodes[0]["date_debut"] is None or periodes[0]["date_debut"] <= DEBUT, banque
        assert periodes[-1]["date_fin"] is None, f"{banque} : la dernière période a une fin"
        for avant, apres in zip(periodes, periodes[1:]):
            assert avant["date_fin"] is not None and apres["date_debut"] is not None, banque
            assert apres["date_debut"] == avant["date_fin"] + dt.timedelta(days=1), \
                f"{banque} : trou ou chevauchement entre {avant['date_fin']} et {apres['date_debut']}"


def test_categories_et_sources():
    """Chaque ligne a une catégorie de la règle, un actionnaire de référence, une source et une page."""
    for l in lignes():
        assert l["actionnariat"] in CATEGORIES, l
        assert l["actionnaire_reference"] and l["source"] and l["page"], l
